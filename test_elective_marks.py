import io
import unittest

import openpyxl

from blueprints.analyzer import calculate_sgpa_and_map_faculty
from marks_utils import is_unselected_subject
from processor import generate_excel_report, update_student_marks
from scraper import parse_vtu_html
from scrapper import parse_vtu_html as parse_legacy_vtu_html


class ElectiveMarkHandlingTests(unittest.TestCase):
    def _parse_subject_rows(self, subject_specs):
        rows = [
            "<tr><td>STUDENT NAME</td><td>Test Student</td></tr>"
        ]
        for code, internal, external in subject_specs:
            total = (
                "-"
                if is_unselected_subject(internal)
                else str(
                    (int(internal) if str(internal).strip().isdigit() else 0)
                    + (int(external) if str(external).strip().isdigit() else 0)
                )
            )
            rows.append(
                f"<tr><td>{code}</td><td>Elective {code}</td>"
                f"<td>{internal}</td><td>{external}</td><td>{total}</td><td>P</td></tr>"
            )
        return parse_vtu_html(
            "2BL23IS001", f"<html><body><table>{''.join(rows)}</table></body></html>"
        )

    def test_internal_dash_controls_subject_applicability(self):
        cases = [
            ("-", "-", False),
            ("-", "50", False),
            ("25", "-", True),
            ("25", "50", True),
            ("0", "-", True),
            ("0", "45", True),
            (" - ", " - ", False),
        ]
        for index, (internal, external, selected) in enumerate(cases):
            with self.subTest(internal=internal, external=external):
                result = self._parse_subject_rows(
                    [(f"21CS{index + 1:02d}", internal, external)]
                )
                subject = next(iter(result["subjects"].values()))
                self.assertEqual(not is_unselected_subject(subject["internal"]), selected)
                if selected:
                    self.assertEqual(subject["total"], int(internal.strip()) + (
                        int(external) if external.isdigit() else 0
                    ))
                else:
                    self.assertEqual(subject["result"], "-")

    def test_zero_and_decimal_internal_marks_are_selected(self):
        for internal in (0, "0", 24, "24", 25.5, "25.5"):
            with self.subTest(internal=internal):
                self.assertFalse(is_unselected_subject(internal))
        for internal in (None, "", "   ", "-", " - ", "–", "—"):
            with self.subTest(internal=internal):
                self.assertTrue(is_unselected_subject(internal))

    def test_parser_aggregates_only_selected_subjects(self):
        result = self._parse_subject_rows([
            ("21CS01", "-", "-"),
            ("21CS02", "-", "50"),
            ("21CS03", "25", "-"),
            ("21CS04", "25", "50"),
            ("21CS05", "0", "-"),
            ("21CS06", "0", "45"),
            ("21CS07", " - ", " - "),
        ])
        self.assertEqual(result["total_marks"], 145)
        self.assertEqual(result["max_marks"], 400)
        self.assertEqual(result["percentage"], 36.25)
        self.assertEqual(result["subjects"]["21CS03"]["external"], "-")
        self.assertEqual(result["subjects"]["21CS03"]["total"], 25)
        self.assertEqual(result["subjects"]["21CS05"]["internal"], 0)

    def test_decimal_internal_with_missing_external_contributes_ia(self):
        result = self._parse_subject_rows([("21CS08", "25.5", "-")])
        self.assertEqual(result["total_marks"], 25.5)
        self.assertEqual(result["max_marks"], 100)
        self.assertEqual(result["subjects"]["21CS08"]["total"], 25.5)

    def test_legacy_parser_uses_the_same_internal_mark_rule(self):
        html = (
            "<table><tr><td>21CS09</td><td>Elective</td>"
            "<td>25</td><td>-</td><td>-</td><td>P</td></tr></table>"
        )
        result = parse_legacy_vtu_html("2BL23IS009", html)
        self.assertEqual(result["total_marks"], 25)
        self.assertEqual(result["max_marks"], 100)
        self.assertEqual(result["subjects"]["21CS09"]["external"], "-")

    def test_students_with_different_electives_get_independent_aggregates(self):
        student_a = self._parse_subject_rows([
            ("21EL01", "25", "50"),
            ("21EL02", "-", "-"),
        ])
        student_b = self._parse_subject_rows([
            ("21EL01", "-", "-"),
            ("21EL02", "28", "48"),
        ])
        calculate_sgpa_and_map_faculty(student_a, [])
        calculate_sgpa_and_map_faculty(student_b, [])

        self.assertEqual((student_a["total_marks"], student_a["max_marks"]), (75, 100))
        self.assertEqual((student_b["total_marks"], student_b["max_marks"]), (76, 100))
        self.assertEqual(student_a["percentage"], 75.0)
        self.assertEqual(student_b["percentage"], 76.0)
        self.assertEqual(student_a["subjects"]["21EL02"]["result"], "-")
        self.assertEqual(student_b["subjects"]["21EL01"]["result"], "-")

    def test_mark_updates_recalculate_selected_subject_totals(self):
        student = {
            "subjects": {
                "SELECTED": {"internal": 25, "external": "-", "total": 25, "result": "F"},
                "NOT_SELECTED": {"internal": "-", "external": "-", "total": 0, "result": "-"},
            }
        }
        update_student_marks(student, "SELECTED", external="-")
        self.assertEqual(student["total_marks"], 25)
        self.assertEqual(student["max_marks"], 100)
        self.assertEqual(student["percentage"], 25.0)
        self.assertEqual(student["status"], "Fail")

    def test_excel_keeps_student_specific_subject_applicability(self):
        students = [
            {
                "usn": "2BL23IS001",
                "name": "Student A",
                "status": "Pass",
                "subjects": {
                    "21EL01": {"internal": 25, "external": 50, "total": 75, "result": "P"},
                    "21EL02": {"internal": "-", "external": "-", "total": 0, "result": "-"},
                },
            },
            {
                "usn": "2BL23IS002",
                "name": "Student B",
                "status": "Pass",
                "subjects": {
                    "21EL01": {"internal": "-", "external": "-", "total": 0, "result": "-"},
                    "21EL02": {"internal": 28, "external": 48, "total": 76, "result": "P"},
                },
            },
        ]
        workbook = openpyxl.load_workbook(
            io.BytesIO(generate_excel_report(students).getvalue()), data_only=False
        )
        sheet = workbook["All Students"]
        self.assertEqual(sheet["D5"].value, 25)
        self.assertEqual(sheet["E5"].value, 50)
        self.assertEqual(sheet["H5"].value, "-")
        self.assertEqual(sheet["I5"].value, "-")
        self.assertIn('IF(D5="-", "-",', sheet["F5"].value)
        self.assertIn('IF(G5<>"-", 1, 0)', sheet["M5"].value)
        self.assertIn("IF(ISNUMBER(INDEX", workbook["Top 5 Marks Secured"]["D4"].value)

    def test_excel_keeps_internal_marks_when_external_is_missing(self):
        student = {
            "usn": "2BL23IS003",
            "name": "Student C",
            "status": "Fail",
            "subjects": {
                "21EL03": {
                    "internal": 25,
                    "external": "-",
                    "total": 25,
                    "result": "F",
                }
            },
        }
        workbook = openpyxl.load_workbook(
            io.BytesIO(generate_excel_report([student]).getvalue()), data_only=False
        )
        sheet = workbook["All Students"]
        self.assertEqual(sheet["D5"].value, 25)
        self.assertEqual(sheet["E5"].value, "-")
        self.assertIn('IF(D5="-", "-",', sheet["F5"].value)

    def test_excel_marks_student_with_no_selected_subjects_as_no_result(self):
        student = {
            "usn": "2BL23IS004",
            "name": "Student D",
            "status": "No Res",
            "subjects": {
                "21EL04": {
                    "internal": "-",
                    "external": "50",
                    "total": 0,
                    "result": "-",
                }
            },
        }
        workbook = openpyxl.load_workbook(
            io.BytesIO(generate_excel_report([student]).getvalue()), data_only=False
        )
        sheet = workbook["All Students"]
        self.assertEqual(sheet["D5"].value, "-")
        self.assertEqual(sheet["E5"].value, "-")
        self.assertIn('"No Res"', sheet["M5"].value)
        self.assertIn('IF((IF(G5<>"-", 1, 0))=0, "-",', sheet["P5"].value)


if __name__ == "__main__":
    unittest.main()
